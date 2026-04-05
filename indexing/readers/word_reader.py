from __future__ import annotations

from pathlib import Path

from indexing.readers.base import BaseReader


def _check_encrypted(path: Path) -> None:
    """Raise ValueError if the Office file is password-protected."""
    try:
        import msoffcrypto

        with open(path, "rb") as f:
            office_file = msoffcrypto.OfficeFile(f)
            if office_file.is_encrypted():
                raise ValueError(f"File is password-protected: {path.name}")
    except ImportError:
        pass  # msoffcrypto not installed — attempt open and let the library report
    except ValueError:
        raise
    except Exception:
        pass  # Not a valid Office container, let downstream library handle it


class WordReader(BaseReader):
    """Reader for Word documents (.docx).

    read_content()  — flat text for Solr fulltext search.
    read_sematic()  — structured text with ### Heading markers and one-row-per-
                      line table format, used by WordChunker for Qdrant embeddings.

    Raises ValueError for password-protected files so the indexing worker
    records a clear error without crashing.
    """

    # Heading styles recognised as section boundaries for semantic chunking.
    _HEADING_STYLES = {"Heading 1", "Heading 2", "Heading 3", "Heading 4",
                       "見出し 1", "見出し 2", "見出し 3", "見出し 4",
                       "Title", "Subtitle"}

    def read_content(self, path: Path) -> str:
        _check_encrypted(path)
        try:
            from docx import Document  # python-docx

            doc = Document(str(path))
            parts: list[str] = []

            for para in doc.paragraphs:
                text = para.text.strip()
                if text:
                    parts.append(text)

            for table in doc.tables:
                for row in table.rows:
                    cells = [c.text.strip() for c in row.cells if c.text.strip()]
                    if cells:
                        parts.append(" | ".join(cells))

            return "\n".join(parts)
        except ValueError:
            raise
        except Exception as exc:
            raise ValueError(f"Could not read Word document '{path.name}': {exc}") from exc

    def read_sematic(self, path: Path) -> str:
        """Return structured text with ### markers for heading-aware chunking."""
        _check_encrypted(path)
        try:
            from docx import Document

            doc = Document(str(path))
            parts: list[str] = []

            # Walk body elements in order to preserve heading → paragraph flow
            for block in doc.element.body:
                tag = block.tag.split("}")[-1] if "}" in block.tag else block.tag

                if tag == "p":
                    from docx.oxml.ns import qn
                    from docx.text.paragraph import Paragraph

                    para = Paragraph(block, doc)
                    text = para.text.strip()
                    if not text:
                        continue
                    style = para.style.name if para.style else ""
                    if style in self._HEADING_STYLES or style.startswith("Heading"):
                        parts.append(f"### {text}")
                    else:
                        parts.append(text)

                elif tag == "tbl":
                    from docx.table import Table

                    table = Table(block, doc)
                    rows_text: list[str] = []
                    header: list[str] = []
                    for i, row in enumerate(table.rows):
                        cells = [c.text.strip() for c in row.cells if c.text.strip()]
                        if i == 0:
                            header = cells
                        else:
                            if header and len(cells) == len(header):
                                # key: value sentence format
                                parts_row = [
                                    f"{h}: {v}" for h, v in zip(header, cells) if v
                                ]
                                if parts_row:
                                    rows_text.append(" / ".join(parts_row))
                            else:
                                if cells:
                                    rows_text.append(" | ".join(cells))
                    if header:
                        rows_text.insert(0, " | ".join(header))
                    parts.extend(rows_text)

            return "\n".join(parts)
        except ValueError:
            raise
        except Exception as exc:
            raise ValueError(f"Could not read Word document '{path.name}': {exc}") from exc
